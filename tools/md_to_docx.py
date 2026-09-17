"""Convert the project's Markdown reports into .docx (and optionally PDF via LibreOffice).

Handles the subset of Markdown the reports use: headings, paragraphs with **bold**,
*italic* and `code`, bullet and numbered lists, pipe tables, images and fenced code.

    python tools/md_to_docx.py docs/DA2_Progress_Report.md [--pdf]
"""

from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x0B, 0x7A, 0x75)
INK = RGBColor(0x1F, 0x29, 0x33)
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_inline(paragraph, text: str, size: float | None = None, color: RGBColor | None = None, bold: bool = False) -> None:
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not token:
            continue
        run_bold = bold
        italic = False
        mono = False
        if token.startswith("**"):
            token, run_bold = token[2:-2], True
        elif token.startswith("`"):
            token, mono = token[1:-1], True
        elif token.startswith("*") and len(token) > 2:
            token, italic = token[1:-1], True
        run = paragraph.add_run(token)
        run.bold = run_bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if mono:
            run.font.name = "Courier New"
            run.font.size = Pt((size or 10.5) - 1)
        if color:
            run.font.color.rgb = color


def convert(md_path: Path) -> Path:
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    for level, size in ((1, 20), (2, 14), (3, 11.5)):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "Cambria"
        st.font.size = Pt(size)
        st.font.color.rgb = TEAL if level == 1 else INK
        st.font.bold = True

    lines = md_path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            block = []
            i += 1
            while not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            i += 1
            continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            i += 1
            continue
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", line)
        if m:
            doc.add_picture(str((md_path.parent / m.group(1)).resolve()), width=Cm(12))
            doc.paragraphs[-1].alignment = 1
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for r, cells in enumerate(rows):
                for c, value in enumerate(cells):
                    cell = table.cell(r, c)
                    cell.paragraphs[0].paragraph_format.space_before = Pt(1.5)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(1.5)
                    if r == 0:
                        shade(cell, "0B7A75")
                        add_inline(cell.paragraphs[0], value, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                    else:
                        add_inline(cell.paragraphs[0], value, size=9.5)
            doc.add_paragraph()
            continue
        m = re.match(r"^(\s*)(- |\d+\. )(.*)", line)
        if m:
            style = "List Bullet" if m.group(2) == "- " else "List Number"
            text = m.group(3)
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and not re.match(r"^\s*(- |\d+\. )", lines[i]):
                text += " " + lines[i].strip()
                i += 1
            add_inline(doc.add_paragraph(style=style), text)
            continue
        text = line.strip()
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|```|\s*- |\s*\d+\. )", lines[i]):
            text += " " + lines[i].strip()
            i += 1
        add_inline(doc.add_paragraph(), text)

    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    src = Path(sys.argv[1])
    docx_path = convert(src)
    print(docx_path)
    if "--pdf" in sys.argv:
        subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)], check=True, capture_output=True)
        print(docx_path.with_suffix(".pdf"))
